"""
File-Level Validation
Pre-parse validation to ensure file integrity before any data extraction
"""
import os
import csv
import chardet
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from dataclasses import dataclass
from enum import Enum
import pandas as pd
import openpyxl


class ValidationStatus(Enum):
    PASSED = "PASSED"
    WARNING = "WARNING"
    REJECTED = "REJECTED"


@dataclass
class ValidationResult:
    status: ValidationStatus
    file_path: str
    messages: List[str]
    metadata: Dict

    def has_errors(self) -> bool:
        return self.status == ValidationStatus.REJECTED

    def has_warnings(self) -> bool:
        return self.status == ValidationStatus.WARNING


class FileValidator:
    """Pre-parse file validation"""

    def __init__(self, strict_mode: bool = True):
        self.strict_mode = strict_mode
        self.max_file_size_mb = 100
        self.min_file_size_bytes = 1

    def validate_file(self, file_path: str) -> ValidationResult:
        """
        Main validation entry point

        Returns:
            ValidationResult with status and messages
        """
        messages = []
        metadata = {}

        # Basic file checks
        if not self._check_file_exists(file_path):
            return ValidationResult(
                status=ValidationStatus.REJECTED,
                file_path=file_path,
                messages=["File does not exist"],
                metadata={}
            )

        # File size check
        file_size = os.path.getsize(file_path)
        if file_size == 0:
            return ValidationResult(
                status=ValidationStatus.REJECTED,
                file_path=file_path,
                messages=["ERROR: File contains 0 bytes"],
                metadata={"size_bytes": 0}
            )

        if file_size > self.max_file_size_mb * 1024 * 1024:
            messages.append(f"WARNING: Large file ({file_size / 1024 / 1024:.1f} MB)")

        metadata['size_bytes'] = file_size

        # Extension validation
        ext = Path(file_path).suffix.lower()
        if ext not in ['.xlsx', '.csv']:
            return ValidationResult(
                status=ValidationStatus.REJECTED,
                file_path=file_path,
                messages=[f"ERROR: Invalid extension '{ext}', expected .xlsx or .csv"],
                metadata=metadata
            )

        # Format-specific validation
        if ext == '.xlsx':
            return self._validate_xlsx(file_path, messages, metadata)
        else:
            return self._validate_csv(file_path, messages, metadata)

    def _check_file_exists(self, file_path: str) -> bool:
        """Check if file exists and is readable"""
        return os.path.isfile(file_path) and os.access(file_path, os.R_OK)

    def _validate_xlsx(self, file_path: str, messages: List[str], metadata: Dict) -> ValidationResult:
        """Excel-specific validation"""
        try:
            wb = openpyxl.load_workbook(file_path, read_only=False, data_only=True)
        except Exception as e:
            return ValidationResult(
                status=ValidationStatus.REJECTED,
                file_path=file_path,
                messages=[f"ERROR: Failed to open Excel file: {str(e)}"],
                metadata=metadata
            )

        # Sheet validation
        sheet_names = wb.sheetnames
        metadata['sheet_names'] = sheet_names
        metadata['sheet_count'] = len(sheet_names)

        if not sheet_names:
            return ValidationResult(
                status=ValidationStatus.REJECTED,
                file_path=file_path,
                messages=["ERROR: No sheets found in workbook"],
                metadata=metadata
            )

        # Check for hidden sheets
        visible_sheets = [name for name in sheet_names if wb[name].sheet_state == 'visible']
        if not visible_sheets:
            return ValidationResult(
                status=ValidationStatus.REJECTED,
                file_path=file_path,
                messages=["ERROR: All sheets are hidden"],
                metadata=metadata
            )

        # Validate first visible sheet
        first_sheet = wb[visible_sheets[0]]

        # Check for merged cells in data region
        merged_cells = self._detect_merged_cells(first_sheet)
        if merged_cells:
            # Allow merges only in first few rows (title area)
            data_region_merges = [mc for mc in merged_cells if mc[0] > 5]  # Row > 5
            if data_region_merges:
                if self.strict_mode:
                    return ValidationResult(
                        status=ValidationStatus.REJECTED,
                        file_path=file_path,
                        messages=[f"ERROR: Merged cells detected in data region: {data_region_merges[:3]}"],
                        metadata=metadata
                    )
                else:
                    messages.append(f"WARNING: Merged cells in data region: {len(data_region_merges)} found")

        # Read data for header detection
        data = []
        for row in first_sheet.iter_rows(min_row=1, max_row=10, values_only=True):
            data.append(row)

        df = pd.DataFrame(data)

        # Detect header row
        header_row = self._detect_header_row(df)
        metadata['header_row'] = header_row

        if header_row > 0:
            messages.append(f"INFO: Header detected in row {header_row + 1}")

        # Detect trailing summaries
        full_data = pd.read_excel(file_path, sheet_name=visible_sheets[0], header=header_row)
        trailing_row = self._detect_trailing_summaries(full_data)
        if trailing_row:
            metadata['trailing_summary_row'] = trailing_row
            messages.append(f"INFO: Trailing summary detected at row {trailing_row}")

        wb.close()

        status = ValidationStatus.WARNING if messages else ValidationStatus.PASSED
        return ValidationResult(
            status=status,
            file_path=file_path,
            messages=messages,
            metadata=metadata
        )

    def _validate_csv(self, file_path: str, messages: List[str], metadata: Dict) -> ValidationResult:
        """CSV-specific validation"""

        # Encoding detection
        encoding = self._detect_encoding(file_path)
        metadata['encoding'] = encoding

        if not encoding:
            return ValidationResult(
                status=ValidationStatus.REJECTED,
                file_path=file_path,
                messages=["ERROR: Failed to detect file encoding"],
                metadata=metadata
            )

        try:
            # Delimiter detection
            with open(file_path, 'r', encoding=encoding) as f:
                sample = f.read(10240)  # Read 10KB sample
                sniffer = csv.Sniffer()

                try:
                    dialect = sniffer.sniff(sample)
                    delimiter = dialect.delimiter
                except Exception:
                    delimiter = ','  # Default

                metadata['delimiter'] = delimiter

            # Column count consistency check
            df = pd.read_csv(file_path, encoding=encoding, delimiter=delimiter, nrows=100)

            # Check for consistent column count
            with open(file_path, 'r', encoding=encoding) as f:
                reader = csv.reader(f, delimiter=delimiter)
                row_lengths = [len(row) for i, row in enumerate(reader) if i < 100]

            if row_lengths:
                expected_cols = row_lengths[0]
                inconsistent_rows = [i for i, length in enumerate(row_lengths) if length != expected_cols]

                if len(inconsistent_rows) > len(row_lengths) * 0.05:  # >5% variance
                    return ValidationResult(
                        status=ValidationStatus.REJECTED,
                        file_path=file_path,
                        messages=[
                            f"ERROR: Inconsistent column counts. Expected {expected_cols}, found variance in {len(inconsistent_rows)} rows"],
                        metadata=metadata
                    )

            metadata['column_count'] = len(df.columns)

        except Exception as e:
            return ValidationResult(
                status=ValidationStatus.REJECTED,
                file_path=file_path,
                messages=[f"ERROR: Failed to parse CSV: {str(e)}"],
                metadata=metadata
            )

        status = ValidationStatus.WARNING if messages else ValidationStatus.PASSED
        return ValidationResult(
            status=status,
            file_path=file_path,
            messages=messages,
            metadata=metadata
        )

    def _detect_encoding(self, file_path: str) -> Optional[str]:
        """Detect file encoding"""
        try:
            with open(file_path, 'rb') as f:
                raw_data = f.read(100000)  # Read 100KB
                result = chardet.detect(raw_data)
                return result['encoding']
        except Exception:
            return None

    def _detect_merged_cells(self, worksheet) -> List[Tuple[int, int]]:
        """Detect merged cells in worksheet"""
        merged = []
        for merged_range in worksheet.merged_cells.ranges:
            # Get top-left cell of merge
            min_row = merged_range.min_row
            min_col = merged_range.min_col
            merged.append((min_row, min_col))
        return merged

    def _detect_header_row(self, df: pd.DataFrame, max_rows: int = 5) -> int:
        """
        Detect most likely header row in first N rows
        Criteria: fewest nulls, most unique values, string-heavy
        """
        scores = []

        for idx in range(min(max_rows, len(df))):
            row = df.iloc[idx]

            # Count non-null values
            non_null = row.notna().sum()

            # Count unique values
            unique = row.nunique()

            # Count string values
            string_count = sum(isinstance(val, str) for val in row)

            # Score: prioritize non-null, unique, strings
            score = non_null * 2 + unique + string_count
            scores.append(score)

        # Return row with highest score
        return scores.index(max(scores)) if scores else 0

    def _detect_trailing_summaries(self, df: pd.DataFrame) -> Optional[int]:
        """Detect trailing summary rows (Total, Average, etc.)"""
        summary_keywords = ['total', 'average', 'grand total', 'sum', 'mean', 'overall']

        # Check last 10 rows
        for idx in range(max(0, len(df) - 10), len(df)):
            row = df.iloc[idx]

            # Check first column for summary keywords
            first_val = str(row.iloc[0]).lower() if len(row) > 0 else ''

            if any(keyword in first_val for keyword in summary_keywords):
                return idx

        return None

    def normalize_header(self, columns: List[str]) -> List[str]:
        """
        Normalize column names

        Rules:
        1. Trim whitespace
        2. Lowercase
        3. Replace non-alphanumeric with underscore
        4. Remove duplicate underscores
        5. Prepend 'col_' if starts with digit
        6. Handle duplicates with _1, _2 suffix
        """
        import re

        normalized = []
        seen: dict = {}

        for col in columns:
            # Basic normalization
            clean = str(col).strip().lower()
            clean = re.sub(r'[^a-z0-9]+', '_', clean)
            clean = re.sub(r'_+', '_', clean)
            clean = clean.strip('_')

            # Prepend if starts with digit
            if clean and clean[0].isdigit():
                clean = f'col_{clean}'

            # Handle empty
            if not clean:
                clean = 'unnamed_column'

            # Handle duplicates
            if clean in seen:
                seen[clean] += 1
                clean = f'{clean}_{seen[clean]}'
            else:
                seen[clean] = 0

            normalized.append(clean)

        return normalized
